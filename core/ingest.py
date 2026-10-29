"""文档解析：把上传的文件 / 内置示例统一变成 LlamaIndex Document。"""

from __future__ import annotations

import io
from pathlib import Path

from llama_index.core import Document

from config import CHUNK_STRATEGIES, SAMPLE_DIR


def _decode(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "gb18030", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        pages.append((page.extract_text() or "").strip())
    return "\n\n".join(p for p in pages if p)


def _read_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return "\n".join(blocks)


def parse_bytes(filename: str, data: bytes) -> str:
    """按扩展名解析成纯文本。不支持的类型直接当文本读。"""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(data)
    if suffix in (".docx", ".doc"):
        if suffix == ".doc":
            raise ValueError("不支持旧版 .doc，请另存为 .docx 或 PDF。")
        return _read_docx(data)
    if suffix in (".txt", ".md", ".markdown", ".csv", ".json", ".log"):
        return _decode(data)
    return _decode(data)


def make_document(filename: str, data: bytes) -> Document:
    text = parse_bytes(filename, data)
    if not text.strip():
        raise ValueError(f"《{filename}》没有解析出任何文字内容。")
    return Document(
        text=text,
        metadata={
            "file_name": filename,
            "file_type": Path(filename).suffix.lower().lstrip("."),
            "char_count": len(text),
        },
        id_=f"doc-{filename}",
    )


def _sample_names(info: dict) -> list[str]:
    """一个策略推荐的内置文档（1~N 份）。兼容旧的单个字符串写法。"""
    raw = info.get("samples")
    if raw is None:
        raw = info.get("sample")
    if isinstance(raw, str):
        return [raw] if raw else []
    return [n for n in raw if n]


def load_samples(keys: list[str]) -> list[Document]:
    """按切片策略 key 载入对应推荐的内置文档。

    同一份文档可能被多个策略推荐（例如最长那份同时属于 Token 和 Markdown），
    多选取时必须按文件名去重，否则会重复入库、检索出两条一模一样的块。
    """
    docs: list[Document] = []
    seen: set[str] = set()
    for key in keys:
        info = CHUNK_STRATEGIES.get(key)
        if not info:
            continue
        for name in _sample_names(info):
            if name in seen:
                continue
            path = SAMPLE_DIR / name
            if path.exists():
                docs.append(make_document(path.name, path.read_bytes()))
                seen.add(name)
    return docs


def sample_catalog() -> list[dict]:
    """内置文档清单（含推荐切片策略与理由），供界面渲染。"""
    catalog = []
    for key, info in CHUNK_STRATEGIES.items():
        names = _sample_names(info)
        paths = [SAMPLE_DIR / n for n in names]
        existing = [p for p in paths if p.exists()]
        catalog.append(
            {
                "strategy": key,
                "filename": names[0] if len(names) == 1 else f"{len(names)} 份文档",
                "filenames": names,
                "exists": bool(existing),
                "size": sum(p.stat().st_size for p in existing),
                "sample_why": info.get("sample_why", ""),
                "fits": info["fits"],
                "why": info["why"],
                "risk": info["risk"],
            }
        )
    return catalog
